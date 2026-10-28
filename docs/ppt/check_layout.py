"""粗略检查 PPT 中文本是否超出文本框/表格单元格（按字符宽度估算换行行数）。
用法：.venv/Scripts/python docs/ppt/check_layout.py"""
import sys
import unicodedata
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

PPTX = Path(__file__).resolve().parents[1] / "GZG语言介绍.pptx"
EMU_PER_PT = 12700


def char_w(c, mono):
    if unicodedata.east_asian_width(c) in ("W", "F"):
        return 1.0
    return 0.55 if mono else 0.52


def para_height(paragraph, width_pt):
    """返回段落估计高度（pt）"""
    runs = paragraph.runs
    if not runs:
        return 12
    size = max((r.font.size.pt if r.font.size else 18) for r in runs)
    mono = any((r.font.name or "").lower() == "consolas" for r in runs)
    text_w = sum(char_w(c, mono) * (r.font.size.pt if r.font.size else 18) for r in runs for c in r.text)
    lines = max(1, int(text_w / max(width_pt, 1)) + 1) if text_w > width_pt else 1
    spacing = paragraph.line_spacing if isinstance(paragraph.line_spacing, float) else 1.15
    return lines * size * 1.2 * spacing


def check_frame(tf, w_emu, h_emu, where, problems, slack=1.08):
    w = w_emu / EMU_PER_PT - 10
    h = h_emu / EMU_PER_PT
    need = sum(para_height(p, w) for p in tf.paragraphs)
    if need > h * slack:
        first = tf.paragraphs[0].text[:30] if tf.paragraphs else ""
        problems.append(f"{where}: 需要约 {need:.0f}pt，可用 {h:.0f}pt  「{first}」")


def main():
    prs = Presentation(PPTX)
    problems = []
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                check_frame(sh.text_frame, sh.width, sh.height, f"第{i}页 文本框", problems)
                if sh.left + sh.width > prs.slide_width + Emu(10000) or sh.top + sh.height > prs.slide_height + Emu(10000):
                    problems.append(f"第{i}页: 形状超出页面边界「{sh.text_frame.text[:20]}」")
            if sh.has_table:
                t = sh.table
                bottom = sh.top + sum(r.height for r in t.rows)
                if bottom > prs.slide_height - Emu(300000):
                    problems.append(f"第{i}页: 表格底部超出页面（{bottom / prs.slide_height:.0%}）")
                for ri, row in enumerate(t.rows):
                    for ci, cell in enumerate(row.cells):
                        check_frame(cell.text_frame, t.columns[ci].width, row.height, f"第{i}页 表格[{ri},{ci}]",
                                    problems, slack=1.6)
    for p in problems:
        print(p)
    print(f"共 {len(prs.slides)} 页，发现 {len(problems)} 处可能溢出")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
