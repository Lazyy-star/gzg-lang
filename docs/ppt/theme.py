"""PPT 与网页共用的配色、字体、代码高亮；以及 PPT 专用的半透明填充、文字与动画 XML 工具。"""
import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

COLORS = {
    "white": "FFFFFF", "text": "E6ECF8", "muted": "93A3C4", "dim": "5E6F96", "ice": "A9CBFF", "sky": "5B95F7",
    "crimson": "E2554F", "warm": "F2A65A", "panel": "050A1E", "kw": "7FB2FF", "str": "F4A07A", "num": "E8777C",
    "cmt": "6D80AA", "call": "9EE6FF", "code": "D8E2F5",
}
SERIF, SANS, MONO = "STZhongsong", "Microsoft YaHei", "Consolas"


def rgb(name):
    return RGBColor.from_string(COLORS.get(name, name))


# ---------------- 代码高亮（两种渲染器共用） ----------------

KEYWORDS = set("let const fn return if elif else while for in break continue struct self match try catch throw import "
               "and or not true false nil 令 常量 函数 返回 如果 否则 否则如果 当 遍历 在 跳出 继续 结构 自身 匹配 尝试 捕获 "
               "抛出 导入 且 或 非 真 假 空".split())
TOKEN = re.compile(r'(#.*$)|("(?:\\.|[^"\\])*"|\'[^\']*\')|(\b\d[\d_.]*\b)|([A-Za-z_\u4e00-\u9fff][\w\u4e00-\u9fff]*)(\s*\()?|(.)')


def highlight(src, lang="gzg"):
    """返回 [[(文本, 颜色名), ...], ...]，每个元素是一行"""
    lines = []
    for line in src.split("\n"):
        out = []
        if lang == "text":
            col = "crimson" if ("错误" in line) else ("warm" if "你是不是想写" in line else "code")
            out.append((line, col))
        elif lang == "git":
            h, _, rest = line.partition(" ")
            m = re.match(r"(\w+)(\W.*)", rest)
            out += [(h + " ", "warm")] + ([(m.group(1), "sky"), (m.group(2), "code")] if m else [(rest, "code")])
        elif lang == "dis":
            body, arrow, note = line.partition("←")
            for part in re.split(r"(\b[A-Z_]{3,}\b|'[^']*'|\(\w+\)|\b\d+\b)", body):
                if not part:
                    continue
                col = "ice" if re.fullmatch(r"[A-Z_]{3,}", part) else "warm" if part.startswith(("'", "(")) else (
                    "dim" if part.strip().isdigit() else "code")
                out.append((part, col))
            if arrow:
                out.append(("←" + note, "crimson"))
        else:
            for m in TOKEN.finditer(line):
                c, s, n, ident, paren, other = m.groups()
                if c:
                    out.append((c, "cmt"))
                elif s:
                    out.append((s, "str"))
                elif n:
                    out.append((n, "num"))
                elif ident:
                    out.append((ident, "kw" if ident in KEYWORDS else "call" if paren else "code"))
                    if paren:
                        out.append((paren, "code"))
                else:
                    out.append((other, "code"))
        lines.append(out)
    return lines


def rich(item, label_color="ice"):
    """'标签::说明' → [(标签, 强调), (说明, 普通)]"""
    if "::" in item:
        a, b = item.split("::", 1)
        return [(a, {"color": label_color, "bold": True}), ("  " + b, {})]
    return [(item, {})]


# ---------------- PPT 工具 ----------------

def set_run(run, size, color="text", bold=False, font=SANS, spacing=None, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    if spacing:
        rpr.set("spc", str(int(spacing * 100)))
    latin = "Georgia" if font == SERIF else font
    for tag, face in (("a:latin", latin), ("a:ea", font if font != MONO else SANS), ("a:cs", latin)):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", face)


def text(slide, x, y, w, h, paras, size=14, color="text", font=SANS, align="l", anchor="t", spacing=None,
         line=1.18, bold=False, space_after=0):
    """paras：字符串，或 [(文本, {color/bold/size/font/italic}), ...] 组成的段落列表"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, para in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = line
        if space_after:
            p.space_after = Pt(space_after)
        for t, st in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            set_run(r, st.get("size", size), st.get("color", color), st.get("bold", bold), st.get("font", font),
                    st.get("spacing", spacing), st.get("italic", False))
    return tb


def _alpha(parent, a):
    clr = parent.find(".//" + qn("a:srgbClr"))
    el = clr.makeelement(qn("a:alpha"), {"val": str(int(a * 100000))})
    clr.append(el)


def shape(slide, x, y, w, h, fill=None, a=1.0, line=None, la=1.0, lw=0.75, kind="round", radius=0.06):
    kinds = {"round": MSO_SHAPE.ROUNDED_RECTANGLE, "rect": MSO_SHAPE.RECTANGLE, "oval": MSO_SHAPE.OVAL,
             "chevron": MSO_SHAPE.CHEVRON}
    s = slide.shapes.add_shape(kinds[kind], x, y, w, h)
    if kind == "round":
        s.adjustments[0] = radius
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        if a < 1:
            _alpha(s._element.spPr.find(qn("a:solidFill")), a)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
        if la < 1:
            _alpha(s._element.spPr.find(qn("a:ln")), la)
    else:
        s.line.fill.background()
    return s
