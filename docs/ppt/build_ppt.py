"""生成《GZG 语言介绍》PPT（深蓝星轨风格，带切换与入场动画）。
  .venv/Scripts/python docs/ppt/make_bg.py     # 先生成背景图（只需一次）
  .venv/Scripts/python docs/ppt/gen_facts.py   # 收集统计数字
  .venv/Scripts/python docs/ppt/build_ppt.py
内容在 content_a/b/c.py（与网页版 build_web.py 共用）。"""
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from anim import animate, transition  # noqa: E402
from blocks import RENDER, ids, r_bullets, r_chart  # noqa: E402
from content_c import all_slides  # noqa: E402
from theme import MONO, SERIF, rich, shape, text  # noqa: E402

ROOT = HERE.parents[1]
OUT = ROOT / "docs" / "GZG语言介绍.pptx"
I = Inches
W, H = 13.333, 7.5
GRID_X, GRID_W = 0.75, 11.83          # 内容网格：12 列 × 6 行
COL = GRID_W / 12


def brand(s, light=True):
    """左上角标识（四格方块 + GZG）与右上角标语"""
    out = []
    for i, (dx, dy) in enumerate([(0, 0), (0.17, 0), (0, 0.17), (0.17, 0.17)]):
        out.append(shape(s, I(0.6 + dx), I(0.42 + dy), I(0.13), I(0.13), fill="FFFFFF" if i != 3 else "E2554F",
                         a=0.95 if i != 3 else 1, radius=0.25))
    out.append(text(s, I(0.98), I(0.36), I(2), I(0.4), "GZG", size=17, color="white", bold=True, font="Segoe UI"))
    out.append(text(s, I(W - 3.6), I(0.36), I(3.0), I(0.5), ["Build a language,", "prove it works."], size=10,
                    color="ice" if light else "muted", align="r", line=1.15))
    return out


def background(s, name):
    return s.shapes.add_picture(str(HERE / "assets" / f"bg_{name}.jpg"), 0, 0, I(W), I(H))


def footer(s, n, total, kicker):
    line = shape(s, I(0.75), I(H - 0.55), I(W - 1.5), Emu(6000), fill="FFFFFF", a=0.15, kind="rect")
    left = text(s, I(0.75), I(H - 0.45), I(6), I(0.3), "GZG · 郭志罡 · 程序设计实践", size=9, color="dim")
    right = text(s, I(W - 2.75), I(H - 0.45), I(2), I(0.3), f"{n:02d} / {total:02d}", size=9, color="dim",
                 align="r", font=MONO)
    return [line, left, right]


def cover(s, d, total):
    bg = background(s, "cover")
    brand(s)
    title = text(s, I(0), I(2.45), I(W), I(1.3), d["title"], size=66, color="white", font=SERIF, align="c", spacing=24)
    en = text(s, I(0), I(3.75), I(W), I(0.5), d["en"], size=16, color="ice", align="c", spacing=3, font="Georgia")
    rule = shape(s, I(W / 2 - 2.6), I(5.95), I(5.2), Emu(9000), fill="FFFFFF", a=0.7, kind="rect")
    sub = text(s, I(0), I(5.3), I(W), I(0.55), d["sub"], size=22, color="white", font=SERIF, align="c", spacing=6)
    author = text(s, I(0), I(6.2), I(W), I(0.4), d["author"], size=13, color="text", align="c", spacing=2)
    animate(s, [ids(title), ids(en), ids(sub, rule), ids(author)], zoom_spid=bg.shape_id, step=420)


def ending(s, d, total):
    bg = background(s, "end")
    brand(s)
    title = text(s, I(0), I(1.2), I(W), I(1.2), d["title"], size=60, color="white", font=SERIF, align="c", spacing=30)
    en = text(s, I(0), I(2.35), I(W), I(0.45), d["en"], size=15, color="ice", align="c", spacing=4, font="Georgia")
    panel = shape(s, I(2.4), I(3.1), I(8.5), I(3.2), fill="050A1E", a=0.45, line="FFFFFF", la=0.18, radius=0.05)
    groups = r_bullets(s, {"items": d["items"], "size": 15}, I(2.8), I(3.4), I(7.8), I(2.8))
    animate(s, [ids(title), ids(en), ids(panel)] + groups, zoom_spid=bg.shape_id)


def content(s, d, n, total, bench):
    background(s, "dark")
    frame = brand(s, light=False) + footer(s, n, total, d.get("kicker", ""))
    kick = text(s, I(0.75), I(0.98), I(8), I(0.3), d.get("kicker", ""), size=11, color="crimson", bold=True,
                spacing=3, font="Segoe UI")
    title = text(s, I(0.75), I(1.25), I(11.8), I(0.75), d["title"], size=32, color="white", font=SERIF, spacing=2)
    head = [kick, title]
    top = 2.12
    if d.get("sub"):
        head.append(text(s, I(0.75), I(2.02), I(11.8), I(0.4), d["sub"], size=14, color="muted"))
        top = 2.55
    row = (H - 0.75 - top) / 6
    groups = [ids(*head)]
    for b in d["blocks"]:
        c, r, cw, rh = b["area"]
        x, y = I(GRID_X + c * COL), I(top + r * row)
        w, h = I(cw * COL - 0.12), I(rh * row - 0.12)
        g = r_chart(s, b, x, y, w, h, bench) if b["type"] == "chart" else RENDER[b["type"]](s, b, x, y, w, h)
        groups += g
    animate(s, groups)
    _ = frame


def main():
    facts = json.loads((HERE / "facts.json").read_text(encoding="utf-8"))
    bench = json.loads((ROOT / "bench" / "results" / "results.json").read_text(encoding="utf-8"))
    slides = all_slides(facts, bench)
    prs = Presentation()
    prs.slide_width, prs.slide_height = I(W), I(H)
    for n, d in enumerate(slides, 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        kind = d.get("kind")
        if kind == "cover":
            cover(s, d, len(slides))
        elif kind == "end":
            ending(s, d, len(slides))
        else:
            content(s, d, n, len(slides), bench)
        transition(s)
    prs.save(OUT)
    print("已生成", OUT, "共", len(prs.slides), "页")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
